# CC-BOM-01 Excel üreticisi: docs/hardware/CC-BOM-01.md ile aynı içerik → docs/hardware/CC-BOM-01.xlsx
# Kullanım (depo kökünden): python3 tools/bom/build_xlsx.py
# BOM değişince önce CC-BOM-01.md, sonra buradaki ROWS güncellenir.
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "docs/hardware/CC-BOM-01.xlsx"
REV = "rev A · 28.10.2026"

G11, G12, G13 = "1.1 Mantık kartı", "1.2 WS2812B şerit", "1.3 Kart montajı"
G2, G3, G4 = "2 Şebeke tarafı", "3 Pano ve kurulum", "4 HIL test"

# (grup, ref, parça, kılıf/anma, görev/not, adet, birim, kapsam, sipariş kalemi; None = parça adı)
ROWS = [
    (G11, "U1", "Mean Well IRM-10-5", "PCB modül, 5 V 2 A 10 W", "AC/DC besleme, çift izolasyonlu (Class II); girişi Sayfa 2 F1 sonrası", 1, "adet", "Zorunlu", None),
    (G11, "U2", "ESP32 DevKit V1 (DOIT, 30 pin)", "ESP32-WROOM-32, 4 MB flash, PSRAM yok", "MCU; kart üstü AMS1117 3.3 V, USB-UART, BOOT (GPIO0), mavi LED (GPIO2)", 1, "adet", "Zorunlu", None),
    (G11, "U3", "2 kanal 5 V röle modülü", "Optokuplörlü, JD-VCC jumper'lı; SRD-05VDC-SL-C, kontak 10 A 250 VAC", "HF (IN1, GPIO32) ve VF (IN2, GPIO33), aktif-LOW. Jumper SÖKÜLÜ: VCC = 3.3 V, JD-VCC = 5 V", 1, "adet", "Zorunlu", None),
    (G11, "B1", "DHT22 / AM2302", "3.3–5 V, tek hat", "T1 + RH1 sensörü, GPIO4", 1, "adet", "Zorunlu", None),
    (G11, "Q1, Q2", "BC337-40", "NPN, TO-92, 45 V 800 mA", "SSR1/SSR2 low-side sürücü", 2, "adet", "Zorunlu", "Transistör BC337-40"),
    (G11, "D1", "SS34", "Schottky 3 A 40 V, SMA", "USB bağlıyken U1'e geri beslemeyi engeller. Delikli kartta eşdeğer: 1N5822", 1, "adet", "Zorunlu", "Diyot SS34 (veya 1N5822)"),
    (G11, "R1", "4.7 kΩ", "¼ W %1", "DHT22 DATA pull-up → 3.3 V", 1, "adet", "Zorunlu", "Direnç 4.7 kΩ ¼ W %1"),
    (G11, "R2", "100 Ω", "¼ W", "DATA seri koruma", 1, "adet", "Zorunlu", "Direnç 100 Ω ¼ W"),
    (G11, "R3, R5", "1 kΩ", "¼ W", "Q1/Q2 taban direnci", 2, "adet", "Zorunlu", "Direnç 1 kΩ ¼ W"),
    (G11, "R4, R6", "10 kΩ", "¼ W", "Taban pull-down: reset/boot anında SSR kapalı", 2, "adet", "Zorunlu", "Direnç 10 kΩ ¼ W"),
    (G11, "R7, R8", "10 kΩ", "¼ W", "Röle IN pull-up → 3.3 V: reset/boot anında röle bırakılmış", 2, "adet", "Zorunlu", "Direnç 10 kΩ ¼ W"),
    (G11, "C1", "470 µF 16 V", "Elektrolitik, 105 °C", "+5V tamponu (D1 çıkışı)", 1, "adet", "Zorunlu", "Kondansatör 470 µF 16 V elektrolitik"),
    (G11, "C2, C4, C5", "100 nF 50 V", "Seramik X7R", "C2 +5V; C4 U3 VCC pinine yakın; C5 sensör ucunda, B1 VCC–GND", 3, "adet", "Zorunlu", "Kondansatör 100 nF 50 V X7R"),
    (G11, "C3", "10 µF 16 V", "Seramik X5R veya tantal", "U3 VCC dekuplajı", 1, "adet", "Zorunlu", "Kondansatör 10 µF 16 V"),
    (G11, "J1", "4 pin klemens / JST-XH", "1 VCC · 2 DATA · 3 NC · 4 GND", "DHT22 kablosu", 1, "adet", "Zorunlu", "Klemens 4 pin"),
    (G11, "J2, J3", "2 pin klemens", "5.08 mm", "SSR1/SSR2 girişi (A1+ / A2−)", 2, "adet", "Zorunlu", "Klemens 2 pin"),
    (G11, "—", "470 Ω", "¼ W", "Yedek: 3.3 V sürüşte röle çekmezse U3'teki 1 kΩ opto direncinin yerine (CC-SCH-01 not 1)", 2, "adet", "İsteğe bağlı", "Direnç 470 Ω ¼ W"),
    (G12, "LS1", "WS2812B şerit, 6 LED", "5 V; 60 LED/m şeritten 6 LED'lik parça (≈ 10 cm) veya 6 modül zinciri", "LED1 durum · LED2 ağ · LED3 MQTT · LED4 mDNS · LED5 ısıtma · LED6 fan", 1, "adet", "Zorunlu", None),
    (G12, "U4", "74AHCT1G125", "SOT-23-5. Delikli kartta eşdeğer: 74HCT125N veya 74AHCT125N (DIP-14); 74HC125 olmaz", "3.3 V → 5 V seviye çevirici; OE → GND, VCC = +5V", 1, "adet", "Önerilen", "Seviye çevirici 74AHCT1G125 (veya 74HCT125N / 74AHCT125N)"),
    (G12, "R13", "330 Ω", "¼ W", "U4 çıkışı (U4 yoksa GPIO27) → R13 → DIN; şeride yakın", 1, "adet", "Zorunlu", "Direnç 330 Ω ¼ W"),
    (G12, "C8", "470 µF 16 V", "Elektrolitik", "Şerit girişinde +5V–GND", 1, "adet", "Zorunlu", "Kondansatör 470 µF 16 V elektrolitik"),
    (G12, "C9", "100 nF 50 V", "Seramik X7R", "Şerit girişinde +5V–GND", 1, "adet", "Zorunlu", "Kondansatör 100 nF 50 V X7R"),
    (G12, "C10", "100 nF 50 V", "Seramik X7R", "U4 VCC dekuplajı (yalnız U4 varsa)", 1, "adet", "Önerilen", "Kondansatör 100 nF 50 V X7R"),
    (G12, "J4", "3 pin klemens / JST-XH", "+5V · DIN · GND", "Şerit bağlantısı", 1, "adet", "Zorunlu", "Klemens 3 pin"),
    (G13, "—", "Delikli pertinaks / prototip kartı", "Çift yüz, ör. 9 × 15 cm", "U1 şebeke pinleri ile SELV arasında ≥ 6 mm açıklık (CC-SCH-01 not 4)", 1, "adet", "Zorunlu", None),
    (G13, "—", "Dişi pin header 1 × 15", "2.54 mm", "DevKit V1 soketi", 2, "adet", "Zorunlu", None),
    (G13, "—", "M3 ara parça + vida/somun", "", "Kart montajı", 4, "takım", "Zorunlu", None),
    (G2, "X1", "Giriş klemensi L / N / PE", "DIN ray, 2.5 mm²; PE yeşil-sarı", "Şebeke girişi (3 kutup)", 1, "takım", "Zorunlu", None),
    (G2, "QF1", "RCBO 1P+N C16, 30 mA, tip A", "DIN ray", "Ana aşırı akım + kaçak akım koruması; sonrası toplam ≈ 9.2 A", 1, "adet", "Zorunlu", None),
    (G2, "F1", "Sigorta T 1 A 5×20 mm + DIN ray yuvası", "250 V", "U1 beslemesi", 1, "adet", "Zorunlu", None),
    (G2, "F2", "MCB 1P C10", "DIN ray", "Rezistans dalı (2 × 4.35 A)", 1, "adet", "Zorunlu", None),
    (G2, "F3", "STB güvenlik termostatı", "Elle resetli, 16 A 250 VAC, ≈ 70 °C; kapiler hazneli veya elle resetli bimetal (KSD301 tipi, 16 A)", "Rezistans dalına seri, yazılımdan bağımsız; algılayıcı ısıtıcı çıkış havasında (CC-SCH-01 not 7)", 1, "adet", "Zorunlu", None),
    (G2, "F4", "Sigorta T 2 A 5×20 mm + DIN ray yuvası", "250 V", "Fan dalı", 1, "adet", "Zorunlu", None),
    (G2, "SSR1, SSR2", "Sıfır geçişli SSR 25 A", "Giriş 3–32 VDC, çıkış ≥ 240 VAC; ör. Crydom D2425, Carlo Gavazzi RM1A23D25", "R1/R2 rezistans anahtarı (SSR_ZC). Ucuz \"SSR-25DA\" kopyalarından kaçının: tipik arıza modu kısa devredir, ARM hattı yoktur", 2, "adet", "Zorunlu", None),
    (G2, "—", "SSR soğutucusu", "Rth ≤ 3 K/W, termal macun veya ped ile", "SSR başına ≈ 5–6 W kayıp (4.35 A); rezistans hava akışının dışında", 2, "adet", "Zorunlu", None),
    (G2, "RV1, RV2", "S14K275 varistör", "275 VAC, 14 mm disk (TDK/EPCOS SIOV)", "SSR çıkışına paralel aşırı gerilim koruması", 2, "adet", "Zorunlu", "Varistör S14K275"),
    (G2, "R11, R12", "100 Ω", "1 W, darbe dayanımlı", "Fan RC snubber direnci", 2, "adet", "Zorunlu", "Direnç 100 Ω 1 W"),
    (G2, "C6, C7", "100 nF X2", "275 VAC, MKP", "Fan RC snubber kondansatörü", 2, "adet", "Zorunlu", "Kondansatör 100 nF X2 275 VAC"),
    (G2, "E1, E2", "Rezistans 1000 W 230 V", "Ör. kanatlı boru rezistans; metal gövde PE terminalli", "R1 / R2 ısıtıcı (4.35 A)", 2, "adet", "Zorunlu", None),
    (G2, "M1", "Fan 230 VAC (Heater Fan)", "Tip ve debi açık konu (OI-H5)", "Rezistans üzerinden üfleme; U3 K1 kontağı", 1, "adet", "Zorunlu", None),
    (G2, "M2", "Fan 230 VAC (Ventilation Fan)", "Tip ve debi açık konu (OI-H5)", "Kulübe havalandırması; U3 K2 kontağı", 1, "adet", "Zorunlu", None),
    (G3, "—", "Pano", "Metal, PE terminalli, en az IP54, DIN ray montaj plakalı; ör. 300 × 400 × 150 mm", "SSR soğutucuları ve mantık kartı için yer", 1, "adet", "Zorunlu", None),
    (G3, "—", "DIN ray TS35", "", "", 1, "m", "Zorunlu", None),
    (G3, "—", "PE barası", "≥ 6 çıkış", "Pano, E1/E2 ve M1/M2 gövdeleri tek baraya yıldız bağlanır", 1, "adet", "Zorunlu", None),
    (G3, "—", "Ray klemensi 2.5 mm²", "Geçiş (gri/mavi) ve PE (yeşil-sarı), köprü elemanlarıyla", "E1, E2, M1, M2 çıkışları (L/N/PE) ve N dağıtımı; ≈ 12–16, yerleşime göre", 16, "adet", "Zorunlu", None),
    (G3, "—", "Kablo rakoru M20", "", "Şebeke, rezistans ve fan kabloları", 5, "adet", "Zorunlu", None),
    (G3, "—", "Kablo rakoru M16", "", "DHT22 kablosu", 1, "adet", "Zorunlu", None),
    (G3, "—", "Kablo 1.5 mm²", "H05VV-F / H07V-K", "Giriş ve rezistans dalı (CC-SCH-01 not 5)", None, "m (g.k.)", "Zorunlu", None),
    (G3, "—", "Kablo 0.75 mm²", "", "Fan dalı", None, "m (g.k.)", "Zorunlu", None),
    (G3, "—", "Kablo 0.25–0.5 mm²", "", "SELV iç bağlantılar", None, "m (g.k.)", "Zorunlu", None),
    (G3, "—", "Sensör kablosu, 4 damarlı ekranlı", "", "3.3 V beslemede ≤ 5 m; daha uzunsa CC-SCH-01 not 3", 5, "m", "Zorunlu", None),
    (G3, "—", "Kablo yüksüğü", "0.5 / 0.75 / 1.5 / 2.5 mm²", "", 1, "takım", "Zorunlu", None),
    (G3, "—", "Yalıtkan bariyer", "", "Şebeke–SELV arasında ≥ 6 mm; SSR giriş/çıkış ve röle kontak tarafı ayrı klemens bloklarında", None, "g.k.", "Zorunlu", None),
    (G3, "—", "Yedek sigorta T 1 A 5×20 mm", "", "", 2, "adet", "Yedek", None),
    (G3, "—", "Yedek sigorta T 2 A 5×20 mm", "", "", 2, "adet", "Yedek", None),
    (G3, "—", "Micro-USB kablo", "", "Programlama ve seri konsol (115200)", 1, "adet", "Zorunlu", None),
    (G4, "R9, R10", "1 kΩ", "¼ W", "HIL-1: J2/J3 yerine +5V → R → LED → Q1/Q2 kollektörü", 2, "adet", "Test", "Direnç 1 kΩ ¼ W"),
    (G4, "LED1, LED2", "3 mm LED", "", "HIL-1 SSR komut göstergesi (WS2812B LED1–LED6 ile karıştırılmasın)", 2, "adet", "Test", None),
    (G4, "—", "Akkor / halojen lamba 25–60 W 230 V + duy", "", "HIL-2 düşük güçlü yük. LED lamba kullanılmaz: SSR kaçak akımı kapalıyken de yakar", 2, "adet", "Test", None),
]

FONT = "Arial"
INK, DIM, WHITE = "1C2733", "586573", "FFFFFF"
HEAD_FILL = PatternFill("solid", fgColor="2F4858")
BAND = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="F1F4F6")]
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="E3E9ED")
WARN_FILL = PatternFill("solid", fgColor="FCE4E1")
THIN = Side(style="thin", color="C9CFD3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical="top", wrap_text=True)
TOP_C = Alignment(vertical="top", horizontal="center", wrap_text=True)
TOP_R = Alignment(vertical="top", horizontal="right")


def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(14, True)
    ws["A2"] = sub
    ws["A2"].font = f(9, color=DIM)


def header(ws, row, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = f(10, True, WHITE)
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def body(c, fill=None, align=TOP, font=None):
    c.font = font or f()
    c.alignment = align
    c.border = BORDER
    if fill:
        c.fill = fill


def fit_height(ws, row, texts_widths, base=13.5):
    lines = 1
    for t, w in texts_widths:
        if t:
            per = max(1, int(w * 1.0))
            lines = max(lines, sum(-(-len(p) // per) or 1 for p in str(t).split("\n")))
    ws.row_dimensions[row].height = base * lines + 3


wb = Workbook()

# ------------------------------------------------------------------ BOM
ws = wb.active
ws.title = "BOM"
title(ws, f"CC-BOM-01 · Kulübe İklim Kontrolörü · Malzeme Listesi · {REV}",
      "Kaynak: docs/hardware/CC-BOM-01.md (CC-SCH-01 rev B + F2.4 WS2812B şeridi, src/app/pins.h). "
      "230 V parçaların kurulumu ve enerjilendirmesi yetkili elektrikçi işidir.")
cols = ["No", "Grup", "Ref.", "Parça / değer", "Kılıf / anma", "Görev / not", "Adet", "Birim", "Kapsam", "Sipariş kalemi"]
widths = [5, 17, 12, 30, 36, 50, 7, 10, 12, 34]
HR = 4
header(ws, HR, cols, widths)
first = HR + 1
groups = []
for i, (g, ref, part, pkg, note, qty, unit, scope, key) in enumerate(ROWS):
    r = first + i
    if g not in groups:
        groups.append(g)
    band = BAND[(len(groups) - 1) % 2]
    font = f(color=DIM, italic=True) if scope == "Test" else f()
    vals = [i + 1, g, ref, part, pkg, note, qty, unit, scope, key or part]
    for col, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=col, value=v)
        body(c, band, TOP_C if col in (1, 7) else TOP, font)
    ws.cell(row=r, column=7).number_format = "0"
    fit_height(ws, r, [(pkg, widths[4]), (note, widths[5]), (part, widths[3]), (key or part, widths[9])])
last = first + len(ROWS) - 1
ws.freeze_panes = ws.cell(row=first, column=4)
ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(cols))}{last}"
ws.cell(row=last + 2, column=2, value="Kapsam: Zorunlu · Önerilen (U4 seviye çevirici ve dekuplajı) · İsteğe bağlı · Yedek · "
        "Test (sahaya gitmez, gri). g.k. = gerektiği kadar; adet boş bırakıldı.").font = f(9, color=DIM)
ws.cell(row=last + 3, column=2, value="Adetler sipariş listesi sayfasına \"Sipariş kalemi\" sütunu üzerinden toplanır; "
        "bir adedi değiştirirseniz sipariş listesi kendiliğinden güncellenir.").font = f(9, color=DIM)

# ------------------------------------------------------------------ Sipariş listesi
so = wb.create_sheet("Sipariş listesi")
title(so, "Sipariş listesi · aynı parçalar birleştirilmiş",
      "Üretim adedi Test satırlarını içermez. Sarı hücreler sizin dolduracağınız alanlardır (fiyat, tedarikçi, durum).")
scols = ["No", "Sipariş kalemi", "Referanslar", "Birim", "Üretim adedi", "Test dahil", "Birim fiyat (₺)", "Tutar (₺)", "Tedarikçi", "Durum"]
swidths = [5, 40, 26, 10, 10, 10, 14, 14, 24, 16]
SH = 4
header(so, SH, scols, swidths)
keys, refs, units = [], {}, {}
for (_, ref, part, _, _, _, unit, _, key) in ROWS:
    k = key or part
    if k not in keys:
        keys.append(k)
        refs[k], units[k] = [], unit
    if ref != "—":
        refs[k].append(ref)
rng = lambda col: f"BOM!${col}${first}:${col}${last}"
s_first = SH + 1
for i, k in enumerate(keys):
    r = s_first + i
    band = BAND[i % 2]
    vals = [
        i + 1, k, ", ".join(refs[k]) or "—", units[k],
        f'=SUMIFS({rng("G")},{rng("J")},B{r},{rng("I")},"<>Test")',
        f"=SUMIFS({rng('G')},{rng('J')},B{r})",
        None, f'=IF(G{r}="","",E{r}*G{r})', None, "Bekliyor",
    ]
    for col, v in enumerate(vals, 1):
        c = so.cell(row=r, column=col, value=v)
        if col in (7, 9, 10):
            body(c, INPUT_FILL, TOP, f(color="0000FF"))
        else:
            body(c, band, TOP_C if col in (1, 5, 6) else TOP)
    so.cell(row=r, column=5).number_format = '0;-0;"–"'
    so.cell(row=r, column=6).number_format = '0;-0;"–"'
    so.cell(row=r, column=7).number_format = '#,##0.00'
    so.cell(row=r, column=8).number_format = '#,##0.00'
    so.cell(row=r, column=8).alignment = TOP_R
    fit_height(so, r, [(k, swidths[1]), (", ".join(refs[k]), swidths[2])])
s_last = s_first + len(keys) - 1
tr = s_last + 1
so.cell(row=tr, column=2, value="Toplam (fiyatı girilen kalemler)")
so.cell(row=tr, column=8, value=f"=SUM(H{s_first}:H{s_last})")
so.cell(row=tr + 1, column=2, value="Fiyatı girilen kalem sayısı")
so.cell(row=tr + 1, column=8, value=f'=COUNT(G{s_first}:G{s_last})&" / "&COUNTA(B{s_first}:B{s_last})')
for rr in (tr, tr + 1):
    for col in range(1, len(scols) + 1):
        body(so.cell(row=rr, column=col), TOTAL_FILL, TOP, f(bold=True))
    so.cell(row=rr, column=8).alignment = TOP_R
so.cell(row=tr, column=8).number_format = '#,##0.00 "₺"'
dv = DataValidation(type="list", formula1='"Bekliyor,Sipariş verildi,Geldi,Stokta var"', allow_blank=True)
dv.error, dv.errorTitle = "Listeden seçin.", "Durum"
so.add_data_validation(dv)
dv.add(f"J{s_first}:J{s_last}")
so.freeze_panes = so.cell(row=s_first, column=3)
so.auto_filter.ref = f"A{SH}:{get_column_letter(len(scols))}{s_last}"
lg = tr + 3
so.cell(row=lg, column=2, value="Nasıl doldurulur").font = f(10, True)
legend = [
    "Sarı hücreler (Birim fiyat, Tedarikçi, Durum) giriş alanıdır; diğerleri formül veya BOM sayfasından gelir.",
    "Birim fiyatı KDV dahil ₺ olarak yazın; Tutar = Üretim adedi × Birim fiyat. Örnek biçim (gerçek fiyat değil): 2,50 → Tutar 2 adet için 5,00.",
    "Tedarikçi serbest metindir (mağaza adı veya bağlantı). Durum listeden seçilir: Bekliyor · Sipariş verildi · Geldi · Stokta var.",
    "Kablo ve bariyer gibi \"g.k.\" kalemlerin adedi boştur (–); uzunluk pano yerleşimine göre belirlenir.",
]
for i, t in enumerate(legend):
    c = so.cell(row=lg + 1 + i, column=2, value=t)
    c.font = f(9, color=DIM)

# ------------------------------------------------------------------ Güç bütçesi
pb = wb.create_sheet("Güç bütçesi")
title(pb, "+5V güç bütçesi · U1 Mean Well IRM-10-5",
      "Değerler datasheet tipik/üst değerlerinden alınmıştır; ölçülmemiştir. Mavi hücreler değiştirilebilir girişlerdir.")
pcols = ["Tüketici", "Adet", "Birim akım (mA)", "Toplam (mA)", "Kaynak / varsayım"]
pwidths = [44, 8, 16, 14, 64]
PH = 4
header(pb, PH, pcols, pwidths)
loads = [
    ("U2 ESP32 DevKit V1 (Wi-Fi gönderim tepesi, AMS1117 üzerinden)", 1, 250, "ESP32 datasheet: 802.11b gönderim ≈ 240 mA; yuvarlandı"),
    ("U3 röle bobini (SRD-05VDC)", 2, 72, "5 V, 0.36 W bobin → ≈ 72 mA"),
    ("SSR girişi", 2, 15, "Üst sınır varsayımı (3–32 VDC girişli SSR, 5 V'ta)"),
    ("LS1 WS2812B LED (beyaz, %100 parlaklık)", 6, 60, "WS2812B: renk başına 20 mA × 3"),
    ("U3 opto, DHT22, pull-up dirençleri", 1, 10, "Tahmini toplam (opto 2 × ≈ 2 mA, DHT22 ≈ 1.5 mA)"),
]
p_first = PH + 1
for i, (name, n, ma, src) in enumerate(loads):
    r = p_first + i
    vals = [name, n, ma, f"=B{r}*C{r}", src]
    for col, v in enumerate(vals, 1):
        c = pb.cell(row=r, column=col, value=v)
        body(c, BAND[i % 2], TOP_C if col in (2, 3, 4) else TOP, f(color="0000FF") if col in (2, 3) else None)
    pb.cell(row=r, column=4).number_format = "#,##0"
p_last = p_first + len(loads) - 1
ws_row = p_first + 3  # WS2812B satırı
R = p_last + 1
rows_calc = [
    ("Toplam, en kötü durum (mA)", f"=SUM(D{p_first}:D{p_last})", "#,##0", None),
    ("U1 anma akımı (mA)", 2000, "#,##0", "IRM-10-5: 5 V 2 A"),
    ("Kullanım oranı", f"=D{R}/D{R+1}", "0.0%", None),
    ("WS2812B parlaklık ayarı", 0.2, "0%", "Firmware varsayılanı %20 (cc_ledstrip.h)"),
    ("Toplam, varsayılan parlaklıkta (mA)", f"=D{R}-D{ws_row}+D{ws_row}*D{R+3}", "#,##0", None),
    ("U1 çıkış gerilimi (V)", 5.0, "0.00", "IRM-10-5 nominal"),
    ("D1 SS34 ileri yön düşümü (V)", 0.4, "0.00", "≈ 0.8 A'de tipik"),
    ("+5V hattı (V)", f"=D{R+5}-D{R+6}", "0.00", None),
    ("AMS1117 3.3 V için yeterli mi (≥ 4.4 V)", f'=IF(D{R+7}>=4.4,"Evet","HAYIR")', None, "≈ 1.1 V düşüm payı; HIL-1'de VIN yük altında ölçülmeli"),
    ("WS2812B besleme aralığında mı (3.5–5.3 V)", f'=IF(AND(D{R+7}>=3.5,D{R+7}<=5.3),"Evet","HAYIR")', None, "WS2812B datasheet"),
    ("Röle çeker mi (SRD-05VDC çekme ≤ 3.75 V)", f'=IF(D{R+7}>=3.75,"Evet","HAYIR")', None, "Bobin nominalinin %75'i"),
]
for i, (label, v, fmt, src) in enumerate(rows_calc):
    r = R + i
    is_input = not (isinstance(v, str) and v.startswith("="))
    for col in range(1, 6):
        body(pb.cell(row=r, column=col), TOTAL_FILL)
    pb.cell(row=r, column=1, value=label).font = f(bold=True)
    c = pb.cell(row=r, column=4, value=v)
    c.font = f(bold=True, color="0000FF" if is_input else INK)
    c.alignment = TOP_R
    if fmt:
        c.number_format = fmt
    if src:
        pb.cell(row=r, column=5, value=src)
pb.freeze_panes = pb.cell(row=p_first, column=1)

# ------------------------------------------------------------------ Notlar
nt = wb.create_sheet("Notlar")
title(nt, "Notlar", f"CC-BOM-01 {REV} · ayrıntı: docs/hardware/CC-BOM-01.md")
nt.column_dimensions["A"].width = 4
nt.column_dimensions["B"].width = 120
notes = [
    ("Şebeke gerilimi", None),
    ("Şebeke tarafı (Grup 2 ve 3) kurulumu ve enerjilendirmesi yetkili elektrikçi tarafından yapılır. RCBO, STB ve PE bağlantıları ölçülmeden enerji verilmez (checklist 10). "
     "ARM hattı olmadığından SSR kısa devre arızasını yalnız F3 (STB) sınırlar.", "warn"),
    ("Çıkış sürücüleri", None),
    ("Isıtıcılar (E1, E2) sıfır geçişli SSR ile sürülür: 20 s'lik zaman-oransal pencerede kanal başına saatte ≈ 180 anahtarlama yapılır; kontaklı röle bu yükte kısa sürede aşınır.", None),
    ("Fanlar (M1, M2) U3 modülündeki kontaklı rölelerle sürülür: seyrek anahtarlanır (min. ON 60 s / OFF 30 s), kontak gerçek açık devre sağlar; RC snubber (R11/C6, R12/C7) kontak arkını ve EMI'yi sınırlar.", None),
    ("Polarite: R1/R2 (GPIO25/26) aktif-HIGH, NPN low-side sürücü, tabanda pull-down. HF/VF (GPIO32/33) aktif-LOW, IN hattında pull-up. Farklı tetik polaritesinde bir SSR veya röle kartı seçilirse src/app/pins.h ve pull dirençleri değişir.", None),
    ("LED'ler", None),
    ("Sahada 6 adet WS2812B LED kullanılır (cc_ledstrip.h, kLedCount = 6). Kart üstü mavi LED (GPIO2) DevKit V1'in üzerindedir, ayrıca alınmaz. HIL LED1/LED2 yalnız ilk testte kullanılır.", None),
    ("CC-SCH-01 rev B ile farklar", None),
    ("1. WS2812B şerit katı (LS1, U4, R13, C8–C10, J4) şemada çizili değil; F2.4'te pins.h'e eklendi (GPIO27). Referanslar bu listede önerilmiştir.", None),
    ("2. D1 (SS34) ve U4 (74AHCT1G125) yüzey montajlıdır; kart delikli pertinaksa 1N5822 ve 74HCT125N (veya 74AHCT125N) kullanılır; T'siz 74HC125 3.3 V girişi güvenilir okumaz.", None),
    ("3. Şema Sayfa 2 PE notundaki \"U1 FG\": IRM-10-5 Class II modüldür, FG pini yoktur.", None),
    ("4. DHT22 kablo sınırı HIL.md'de ≤ 20 m, şema notunda (3.3 V beslemede) ≤ 5 m; bu listede şema notu esas alınmıştır.", None),
    ("v1 kapsamı dışında (pinleri ayrılmış, bu listede yok)", None),
    ("HEATER_ARM hattı — GPIO13, lojik VE / ikinci transistör katı (OI-H6, ADR-008) · T2 hava çıkış sensörü — GPIO18 1-Wire, DS18B20 + 4.7 kΩ (OI-H2) · "
     "RTC — GPIO21/22 I²C, DS3231 (OI-H10; v1 NTP kullanır) · Fan geri bildirimi — akım trafosu veya RPM çıkışlı fan (OI-H7).", None),
]
r = 4
for text, kind in notes:
    c = nt.cell(row=r, column=2, value=text)
    if kind is None and len(text) < 60:
        c.font = f(11, True)
        r += 1
        continue
    c.font = f()
    c.alignment = TOP
    if kind == "warn":
        c.fill = WARN_FILL
        c.font = f(bold=True, color="8A1C14")
    fit_height(nt, r, [(text, 120)])
    r += 1

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
    sh.page_setup.orientation = "landscape"
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    if sh.title != "Notlar":
        sh.print_title_rows = "4:4"
ws.cell(row=HR, column=10).comment = Comment("Sipariş listesi bu sütuna göre gruplanır (SUMIFS).", "CC-BOM-01")
wb.save(OUT)
print(OUT, len(ROWS), "satır,", len(keys), "sipariş kalemi")
